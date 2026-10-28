#!/usr/bin/env python
"""L7 QA: make the RA packet a pure compare-and-click task. Run after 91_qa_triage.py.

For every row in data/interim/telegram/qa/ra_packet/ra_{impacts,empty,typing}.csv:
  - machine-translates the post to English (text_en; model llm.MODEL_EXTRACT, inference only, cached on disk);
  - renders the AI's extracted fields in English (place_en, weapon_en) with the original in brackets;
  - turns "how many strikes did the AI miss" into a yes/no (chk_all_found) plus an optional count;
  - writes the CSVs back (same names, so 92_qa_score.py and qa_checker.html read them) and an .xlsx copy of each with
    wrapped text and Yes/No dropdowns for RAs who prefer a spreadsheet.
The Ukrainian original stays in every sheet (column `text`); the guide tells RAs to check it when a detail decides
the answer. Translation is labelled machine-made everywhere.

  .venv/bin/python -I scripts/telegram/93_qa_translate.py            # translate + rewrite packet
  .venv/bin/python -I scripts/telegram/93_qa_translate.py --dry-run  # counts only, no network
"""
from __future__ import annotations
import argparse, sys, pathlib
from concurrent.futures import ThreadPoolExecutor
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import pandas as pd
import llm
from _common import INTERIM, write_meta, log

PACKET = INTERIM / "qa" / "ra_packet"
PV = "translate_v1"
MODEL = llm.MODEL_EXTRACT
SYSTEM = """You translate Ukrainian (sometimes Russian) Telegram posts from Ukrainian regional administrations into plain English.
Translate faithfully and completely: keep every place name, number, time, weapon and casualty detail exactly; do not
summarise, add or explain. Transliterate place names in the usual English form and keep the original in brackets the
first time, e.g. "Kryvyi Rih (Кривий Ріг)". Keep emojis out. Also translate each item of `terms` (place names: give the
usual English transliteration; weapon words: give a short English term) in the same order."""
SCHEMA = {"type": "object", "additionalProperties": False, "required": ["text_en", "terms_en"],
          "properties": {"text_en": {"type": "string"}, "terms_en": {"type": "array", "items": {"type": "string"}}}}

IMP_ORDER = ["row_id", "text_en", "place_en", "weapon_en", "target_type", "killed", "wounded", "event_date", "event_time",
             "n_incidents_in_post", "is_debris_only", "location_withheld",
             "chk_place_ok", "chk_weapon_ok", "chk_target_type_ok", "chk_casualties_ok", "chk_time_ok",
             "chk_all_found", "chk_missed_incidents_in_post", "chk_note", "text", "url"]
EMP_ORDER = ["row_id", "text_en", "chk_has_incident", "chk_note", "text", "url"]
TYP_ORDER = ["row_id", "text_en", "post_type", "human_ok", "human_post_type", "human_note", "text", "url"]
YN = {"imp": ["chk_place_ok", "chk_weapon_ok", "chk_target_type_ok", "chk_casualties_ok", "chk_time_ok", "chk_all_found"],
      "emp": ["chk_has_incident"], "typ": ["human_ok"]}


def s(v) -> str:
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()


def place_terms(r) -> list[str]:
    return [s(getattr(r, c, "")) for c in ("settlement_raw", "raion_raw", "oblast_raw")]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--workers", type=int, default=6)
    a = ap.parse_args()

    imp = pd.read_csv(PACKET / "ra_impacts.csv", dtype=str)
    emp = pd.read_csv(PACKET / "ra_empty.csv", dtype=str)
    typ = pd.read_csv(PACKET / "ra_typing.csv", dtype=str)
    if "weapon_raw" not in imp.columns:
        sys.exit("packet already translated; re-run 91_qa_triage.py first (cached, free), then this script")

    # one request per unique post; terms = every place/weapon string any row of that post needs
    jobs: dict[str, list[str]] = {}
    for df, has_terms in ((imp, True), (emp, False), (typ, False)):
        for r in df.itertuples():
            t = jobs.setdefault(s(r.text), [])
            if has_terms:
                for x in place_terms(r) + [s(r.weapon_raw)]:
                    if x and x not in t:
                        t.append(x)
    jobs.pop("", None)
    log(f"{len(jobs)} unique posts to translate ({MODEL}); rows: impacts {len(imp)}, empty {len(emp)}, typing {len(typ)}")
    if a.dry_run:
        return

    def one(item):
        text, terms = item
        user = f"POST:\n{text}\n\nterms: {terms}"
        try:
            out = llm.call(PV, SYSTEM, user, SCHEMA, MODEL, max_tokens=4096, effort="low")
            tr = dict(zip(terms, out.get("terms_en") or []))
            return text, out["text_en"].strip(), tr
        except Exception as e:
            log(f"translation failed: {type(e).__name__}"); return text, None, {}
    with ThreadPoolExecutor(a.workers) as ex:
        res = {t: (en, tr) for t, en, tr in ex.map(one, jobs.items())}
    n_fail = sum(v[0] is None for v in res.values())
    if n_fail:
        sys.exit(f"{n_fail} posts failed to translate; packet NOT rewritten. Re-run (successes are cached).")

    def en(text):
        return "[Machine translation, check the Ukrainian original if a detail decides your answer]\n\n" + res[s(text)][0]

    def term(text, x):
        x = s(x)
        if not x:
            return ""
        e = res[s(text)][1].get(x, "")
        return f"{e} ({x})" if e and e.casefold() != x.casefold() else x

    imp["text_en"] = imp.text.map(en)
    imp["place_en"] = [", ".join(p for p in (term(r.text, r.settlement_raw), term(r.text, r.raion_raw), term(r.text, r.oblast_raw)) if p)
                       or ("withheld by the post" if s(r.location_withheld).lower() == "true" else "(none)") for r in imp.itertuples()]
    imp["weapon_en"] = [term(r.text, r.weapon_raw) or "(none named)" for r in imp.itertuples()]
    imp["chk_all_found"] = ""
    for c in ("killed", "wounded"):
        imp[c] = imp[c].fillna("").map(lambda v: "none reported" if s(v) in ("", "0", "0.0") else str(int(float(v))))
    emp["text_en"] = emp.text.map(en)
    typ["text_en"] = typ.text.map(en)

    sheets = {"ra_impacts": (imp, IMP_ORDER, "imp"), "ra_empty": (emp, EMP_ORDER, "emp"), "ra_typing": (typ, TYP_ORDER, "typ")}
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    for name, (df, order, k) in sheets.items():
        df = df[[c for c in order if c in df.columns]]
        df.to_csv(PACKET / f"{name}.csv", index=False)
        wb = Workbook(); ws = wb.active; ws.title = name
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([s(v) for v in row])
        wide = {"text_en": 70, "text": 50, "place_en": 32, "weapon_en": 22, "chk_note": 30, "human_note": 30, "url": 28}
        for j, c in enumerate(df.columns, 1):
            col = ws.cell(row=1, column=j).column_letter
            ws.column_dimensions[col].width = wide.get(c, 14)
            ws.cell(row=1, column=j).font = Font(bold=True)
            if c.startswith(("chk_", "human_")):
                ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor="FFF2CC")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"y,n"', allow_blank=True)
        ws.add_data_validation(dv)
        for c in YN[k]:
            j = list(df.columns).index(c) + 1
            col = ws.cell(row=1, column=j).column_letter
            dv.add(f"{col}2:{col}{len(df) + 1}")
        ws.freeze_panes = "B2"
        wb.save(PACKET / f"{name}.xlsx")
    write_meta(PACKET / "ra_impacts.csv", inputs={"ra_impacts": PACKET / "ra_impacts.csv"},
               counts={"posts_translated": len(res), "impacts": len(imp), "empty": len(emp), "typing": len(typ)},
               prompt_version=PV)
    log(f"packet rewritten with English in {PACKET} (.csv for the checker page, .xlsx for spreadsheet users)")


if __name__ == "__main__":
    main()
